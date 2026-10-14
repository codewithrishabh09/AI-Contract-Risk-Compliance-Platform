from io import BytesIO

from docx import Document

from app.services.extraction.models import (
    ExtractedDocument,
    ExtractedPage,
)
from app.services.extraction.normalizer import normalize_text


class DOCXExtractionError(Exception):
    """Raised when DOCX extraction fails."""


PARAGRAPHS_PER_LOGICAL_PAGE = 40


def extract_docx(
    *,
    content: bytes,
    filename: str,
    content_type: str,
) -> ExtractedDocument:
    """Extract DOCX paragraphs into logical page groups."""

    try:
        document = Document(
            BytesIO(content),
        )
    except Exception as exc:
        raise DOCXExtractionError(
            "Unable to open the DOCX document."
        ) from exc

    paragraphs: list[str] = []

    for paragraph in document.paragraphs:
        text = normalize_text(
            paragraph.text,
        )

        if text:
            paragraphs.append(text)

    pages: list[ExtractedPage] = []

    for start in range(
        0,
        len(paragraphs),
        PARAGRAPHS_PER_LOGICAL_PAGE,
    ):
        group = paragraphs[
            start:start + PARAGRAPHS_PER_LOGICAL_PAGE
        ]

        page_text = "\n".join(group)

        pages.append(
            ExtractedPage(
                page_number=(start // PARAGRAPHS_PER_LOGICAL_PAGE) + 1,
                text=page_text,
                character_count=len(page_text),
            )
        )

    return ExtractedDocument(
        filename=filename,
        content_type=content_type,
        pages=pages,
    )